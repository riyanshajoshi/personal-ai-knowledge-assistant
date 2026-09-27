import os
from pypdf import PdfReader
from pptx import Presentation

def load_pdf(filepath):
    """Extract all text from a PDF file."""
    reader = PdfReader(filepath)
    text = ""
    for page in reader.pages:
        text += page.extract_text() + "\n"
    return text

def load_pptx(filepath):
    """Extract all text from a PowerPoint file."""
    prs = Presentation(filepath)
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        text += run.text + " "
                text += "\n"
    return text

def load_text(filepath):
    """Read a plain text or markdown file."""
    with open(filepath, "r", encoding="utf-8") as f:
        return f.read()

def load_documents(folder_path):
    """Walk a folder and load all supported files into a list of {source, text} dicts."""
    documents = []
    for filename in os.listdir(folder_path):
        filepath = os.path.join(folder_path, filename)
        if filename.lower().endswith(".pdf"):
            text = load_pdf(filepath)
        elif filename.lower().endswith(".pptx"):
            text = load_pptx(filepath)
        elif filename.lower().endswith((".txt", ".md")):
            text = load_text(filepath)
        else:
            continue  # skip unsupported files
        documents.append({"source": filename, "text": text})
    return documents

if __name__ == "__main__":
    docs = load_documents("my_knowledge_base")
    print(f"Loaded {len(docs)} documents:")
    for d in docs:
        print(f" - {d['source']} ({len(d['text'])} characters)")